import os
import fitz
import uuid
from cnocr import CnOcr
from PIL import Image
import logging
from docx import Document
from multiprocessing.pool import ThreadPool
from functools import partial
from tqdm import tqdm
import multiprocessing as mp
import openpyxl
import xlrd

logging.basicConfig(filename='extraction.log', level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

INPUT_DIR = "[PATH_TO_YOUR_INPUT_DOCUMENTS]"
OUTPUT_DIR = "[PATH_TO_YOUR_TEXT_OUTPUT]"
TEMP_IMG_DIR = "[PATH_TO_TEMP_IMAGES]"
os.makedirs(TEMP_IMG_DIR, exist_ok=True)
os.makedirs(OUTPUT_DIR, exist_ok=True)

ocr_model = CnOcr()

def is_image_based_page(page):
    text = page.get_text().strip()
    return len(text) < 10

def extract_text_from_pdf_page(page):
    try:
        return page.get_text("text").strip()
    except Exception as e:
        logging.error(f"Error extracting text from PDF page: {e}")
        return ""

def extract_text_from_cnocr(page, page_index):
    try:
        img_filename = f"page_{uuid.uuid4().hex[:8]}.jpeg"
        img_path = os.path.join(TEMP_IMG_DIR, img_filename)

        pix = page.get_pixmap(dpi=200)
        pix.save(img_path)

        result = ocr_model.ocr(img_path)
        text = "\n".join([item['text'] for item in result])

        os.remove(img_path)
        return text
    except Exception as e:
        logging.error(f"CnOcr error on page {page_index + 1}: {e}")
        return ""

def process_pdf(file_path, output_dir):
    rel_path = os.path.relpath(file_path, INPUT_DIR)
    output_subdir = os.path.join(output_dir, os.path.dirname(rel_path))
    os.makedirs(output_subdir, exist_ok=True)
    filename = os.path.splitext(os.path.basename(file_path))[0]
    output_path = os.path.join(output_subdir, f"{filename}_extracted.txt")

    if os.path.exists(output_path):
        logging.info(f"Skipping {file_path}: Output already exists")
        print(f"Skipping {file_path}: Output already exists")
        return

    logging.info(f"Processing {file_path}...")
    print(f"Processing {file_path}...")
    page_buffer = ""

    try:
        doc = fitz.open(file_path)
        for i, page in enumerate(doc):
            page_text = f"\n--- Page {i + 1} ---\n"
            if is_image_based_page(page):
                logging.info(f"Page {i + 1} is image-based, using CnOcr...")
                page_text += extract_text_from_cnocr(page, i)
            else:
                logging.info(f"Page {i + 1} is text-based, extracting text...")
                page_text += extract_text_from_pdf_page(page)

            page_buffer += page_text

            if (i + 1) % 5 == 0 or (i + 1) == len(doc):
                with open(output_path, 'a', encoding='utf-8') as f:
                    f.write(page_buffer)
                page_buffer = ""

        doc.close()
    except Exception as e:
        logging.error(f"Error processing PDF {file_path}: {e}")
        print(f"Error processing PDF {file_path}: {e}")

def process_docx(file_path, output_dir):
    rel_path = os.path.relpath(file_path, INPUT_DIR)
    output_subdir = os.path.join(output_dir, os.path.dirname(rel_path))
    os.makedirs(output_subdir, exist_ok=True)
    filename = os.path.splitext(os.path.basename(file_path))[0]
    output_path = os.path.join(output_subdir, f"{filename}_extracted.txt")

    if os.path.exists(output_path):
        logging.info(f"Skipping {file_path}: Output already exists")
        return

    try:
        doc = Document(file_path)
        outline_text = []
        seen_content = set()

        for table_index, table in enumerate(doc.tables):
            logging.info(f"{file_path}: Processing table {table_index + 1} with {len(table.rows)} rows and {len(table.columns)} columns")
            for row_index, row in enumerate(table.rows):
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if not cells:
                    continue

                unique_cells = list(dict.fromkeys(cells))
                merged_cells = []
                current_text = ""
                for cell in unique_cells:
                    if len(cell) == 1:
                        current_text += cell
                    else:
                        if current_text:
                            merged_cells.append(current_text)
                            current_text = ""
                        merged_cells.append(cell)
                if current_text:
                    merged_cells.append(current_text)

                content_key = tuple(merged_cells)
                if content_key in seen_content:
                    continue
                seen_content.add(content_key)

                text = merged_cells[0] if merged_cells else ""
                level = 0
                if text.startswith(('1.', '2.', '3.', '4.', '5.')):
                    level = 0
                elif text.startswith(('1.1', '1.2', '2.1', '2.2')):
                    level = 1
                elif text.startswith(('a.', 'b.', 'c.', 'i.', 'ii.')):
                    level = 2
                elif text.startswith(('i)', 'ii)')):
                    level = 3

                content = " ".join(merged_cells[1:]) if len(merged_cells) > 1 else text
                indent = "  " * level
                outline_text.append(f"{indent}{content}")

        result = "\n".join(outline_text) if outline_text else "No table data extracted."

        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(result)
        logging.info(f"Saved extracted outline to {output_path}")
    except Exception as e:
        logging.error(f"Error processing DOCX file {file_path}: {e}")

def process_excel(file_path, output_dir):
    rel_path = os.path.relpath(file_path, INPUT_DIR)
    output_subdir = os.path.join(output_dir, os.path.dirname(rel_path))
    os.makedirs(output_subdir, exist_ok=True)
    filename = os.path.splitext(os.path.basename(file_path))[0]
    output_path = os.path.join(output_subdir, f"{filename}_extracted.txt")

    if os.path.exists(output_path):
        logging.info(f"Skipping {file_path}: Output already exists")
        return

    extracted_text = ""
    try:
        if file_path.endswith(".xlsx"):
            wb = openpyxl.load_workbook(file_path, read_only=True)
            for sheet in wb.worksheets:
                extracted_text += f"\n--- Sheet: {sheet.title} ---\n"
                for row in sheet.iter_rows(values_only=True):
                    row_text = "\t".join([str(cell) if cell is not None else "" for cell in row])
                    extracted_text += row_text + "\n"
        elif file_path.endswith(".xls"):
            wb = xlrd.open_workbook(file_path)
            for sheet in wb.sheets():
                extracted_text += f"\n--- Sheet: {sheet.name} ---\n"
                for row_idx in range(sheet.nrows):
                    row_values = sheet.row_values(row_idx)
                    row_text = "\t".join([str(cell) for cell in row_values])
                    extracted_text += row_text + "\n"

        if extracted_text.strip():
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(extracted_text)
            logging.info(f"Saved extracted Excel to {output_path}")
    except Exception as e:
        logging.error(f"Error processing Excel file {file_path}: {e}")

def get_all_documents(root_dir):
    files = []
    for root, _, filenames in os.walk(root_dir):
        for file in filenames:
            if file.lower().endswith(('.pdf', '.docx', '.xlsx', '.xls')):
                files.append(os.path.join(root, file))
    return files

def main():
    files = get_all_documents(INPUT_DIR)
    pdf_files = [f for f in files if f.lower().endswith('.pdf')]
    docx_files = [f for f in files if f.lower().endswith('.docx')]
    excel_files = [f for f in files if f.lower().endswith(('.xlsx', '.xls'))]

    pool = ThreadPool(processes=2)
    process_func = partial(process_pdf, output_dir=OUTPUT_DIR)
    for _ in tqdm(pool.imap_unordered(process_func, pdf_files), total=len(pdf_files), desc="Processing PDF files"):
        pass
    pool.close()
    pool.join()

    for file_path in tqdm(docx_files, desc="Processing DOCX files"):
        process_docx(file_path, OUTPUT_DIR)

    for file_path in tqdm(excel_files, desc="Processing Excel files"):
        process_excel(file_path, OUTPUT_DIR)

    logging.info("All files processed.")
    print("All files processed.")

if __name__ == "__main__":
    mp.set_start_method('spawn', force=True)
    main()